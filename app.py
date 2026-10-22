from io import BytesIO
from decimal import Decimal
import os
from pathlib import Path
from xml.sax.saxutils import escape
from flask import Flask, render_template, request, jsonify, send_file
from werkzeug.exceptions import RequestEntityTooLarge
from comparator import extract_pdf, compare, supplier_label

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024
APP_VERSION = '0.6.1'

@app.after_request
def disable_stale_pages(response):
    if request.endpoint != 'static':
        response.headers['Cache-Control'] = 'no-store'
    return response

@app.get('/')
def home():
    return render_template('index.html', version=APP_VERSION)

@app.get('/versao')
def version():
    return jsonify(versao=APP_VERSION, instancia=os.environ.get('COMPARADOR_INSTANCE_ID', ''))

@app.get('/exemplos/<name>')
def example(name):
    if name not in ('sped_simulado.pdf', 'dominio_simulado.pdf'):
        return jsonify(error='Exemplo não encontrado.'), 404
    return send_file(Path(app.root_path) / 'exemplos' / name,
                     mimetype='application/pdf', as_attachment=True)

@app.errorhandler(RequestEntityTooLarge)
def too_large(error):
    return jsonify(error='O envio deve ter no máximo 20 MB no total.'), 413

class ComparisonError(ValueError):
    def __init__(self, message, sources):
        super().__init__(message)
        self.sources = sources


def read_source(file):
    if not file or not file.filename or not file.filename.lower().endswith('.pdf'):
        raise ValueError('Selecione um arquivo PDF.')
    if file.stream.read(5) != b'%PDF-':
        raise ValueError('Arquivo inválido: o conteúdo deve ser PDF.')
    file.stream.seek(0)
    metadata = {}
    try:
        notes, issues = extract_pdf(file.stream, metadata=metadata)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Não foi possível ler o PDF. Verifique se está íntegro e sem senha.') from exc
    summary = dict(quantidade=len(notes), paginas=metadata['paginas'],
                   total=str(sum((Decimal(n['valor']) for n in notes), Decimal('0.00'))),
                   cnpj=sum(n['tipo_identificacao'] == 'CNPJ' for n in notes),
                   cpf=sum(n['tipo_identificacao'] == 'CPF' for n in notes),
                   warnings=issues, error=None)
    if 'total_declarado' in metadata:
        summary['total_relatorio'] = metadata['total_declarado']
        summary['total_confere'] = Decimal(summary['total']) == Decimal(metadata['total_declarado'])
    if not notes:
        summary['error'] = 'Nenhuma nota reconhecida. O layout precisa de ajuste com o PDF original; se não houver texto extraível, será necessário OCR.'
    return notes, summary


def failure(exc):
    payload = dict(error=str(exc))
    if isinstance(exc, ComparisonError):
        payload['sources'] = exc.sources
    return jsonify(payload), 422


@app.post('/conferir/<source>')
def inspect_source(source):
    if source not in ('sped', 'dominio'):
        return jsonify(error='Origem inválida.'), 404
    try:
        notes, summary = read_source(request.files.get(source))
    except ValueError as exc:
        return failure(exc)
    response = dict(source=source, notes=notes, summary=summary)
    if summary['error']:
        response['error'] = summary['error']
        return jsonify(response), 422
    return jsonify(response)


def analyze():
    notes, summaries, warnings = [], {}, []
    for name in ('sped', 'dominio'):
        try:
            found, summary = read_source(request.files.get(name))
        except ValueError as exc:
            found, summary = [], dict(quantidade=0, error=str(exc), warnings=[])
        notes.append(found)
        summaries[name] = summary
        warnings.extend(f'{name.upper()}: {w}' for w in summary['warnings'])
    if any(summary['error'] for summary in summaries.values()):
        errors = '; '.join(f'{name.upper()}: {summary["error"]}' for name, summary in summaries.items() if summary['error'])
        raise ComparisonError(errors + ' Use “Conferir leitura” para revisar cada PDF separadamente.', summaries)
    return dict(rows=compare(*notes), warnings=warnings, counts=[len(s) for s in notes], sources=summaries)

@app.post('/comparar')
def comparison():
    try:
        return jsonify(analyze())
    except ValueError as exc:
        return failure(exc)

@app.post('/exportar/<kind>')
def export(kind):
    if kind not in ('xlsx', 'pdf'): return jsonify(error='Formato inválido.'), 404
    try:
        result = analyze()
    except ValueError as exc:
        return failure(exc)
    headers = ['Situação', 'CNPJ/CPF', 'Documento', 'Diferenças', 'Fornecedor SPED', 'Valor SPED', 'Fornecedor Domínio', 'Valor Domínio', 'Pendências']
    rows = []
    for r in result['rows']:
        # Keep every duplicate visible rather than silently choosing one.
        rows.append([r['status'], r['cnpj'], r['documento'], r['diferencas']] + [ '\n'.join(supplier_label(n) if f == 'fornecedor' else n[f] for n in r[s]) for s in ('sped','dominio') for f in ('fornecedor','valor')] + [' '.join(r['pendencias'])])
    output = BytesIO()
    if kind == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        wb = Workbook(); ws = wb.active; ws.title = 'Comparativo'
        for row in [headers] + rows:
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in ws[1]: cell.font = Font(bold=True)
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 25
        info = wb.create_sheet('Avisos'); info.append(['Conferir extração antes de usar o relatório.'])
        for warning in result['warnings']: info.append([warning])
        wb.save(output)
        mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.pagesizes import A4, landscape
        styles = getSampleStyleSheet(); style = styles['BodyText']; style.fontSize = 7; style.leading = 9
        story = [Paragraph('Comparativo de notas de entrada — SPED x Domínio', styles['Title']), Paragraph('Confira a extração antes de utilizar os resultados.', style), Spacer(1,12)]
        table = Table([[Paragraph(escape(str(c)).replace('\n', '<br/>'), style) for c in row] for row in [headers]+rows], repeatRows=1, colWidths=[75,90,50,55,110,55,110,55,100])
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),'#e4eee9'),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),0.3,'#bbbbbb')]))
        story.append(table)
        for warning in result['warnings']: story.append(Paragraph(escape(warning),style))
        SimpleDocTemplate(output, pagesize=landscape(A4),rightMargin=20,leftMargin=20).build(story)
        mime = 'application/pdf'
    output.seek(0)
    return send_file(output, mimetype=mime, as_attachment=True, download_name=f'comparativo.{kind}')

if __name__ == '__main__':
    app.run(host=os.environ.get('COMPARADOR_HOST', '0.0.0.0'),
            port=int(os.environ.get('COMPARADOR_PORT', '5000')), debug=False)
