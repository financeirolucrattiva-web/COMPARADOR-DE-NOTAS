"""Valida a simulação contra o servidor local e guarda relatórios efetivos."""
from collections import Counter
from io import BytesIO
from pathlib import Path
import json
import urllib.request
import pdfplumber
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

def main():
    examples = ROOT/'exemplos'
    expected = json.loads((examples/'gabarito.json').read_text(encoding='utf-8'))
    boundary = 'LucrattivaSyntheticValidationBoundary'
    body = b''
    for source in ('sped','dominio'):
        body += f'--{boundary}\r\nContent-Disposition: form-data; name="{source}"; filename="{source}.pdf"\r\nContent-Type: application/pdf\r\n\r\n'.encode()
        body += (examples/f'{source}_simulado.pdf').read_bytes() + b'\r\n'
    body += f'--{boundary}--\r\n'.encode()
    opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
    def post(path):
        req=urllib.request.Request('http://127.0.0.1:5000'+path,data=body,headers={'Content-Type':'multipart/form-data; boundary='+boundary})
        with opener.open(req,timeout=30) as response:
            assert response.status==200
            return response.read()
    with opener.open('http://127.0.0.1:5000', timeout=10) as response:
        assert 'Suas notas' in response.read().decode()
    result=json.loads(post('/comparar'))
    assert result['counts']==expected['counts']
    assert not result['warnings']
    assert Counter(r['status'] for r in result['rows'])==Counter(expected['summary'])
    assert len(result['rows'])==105
    for row,target in zip(result['rows'],expected['rows'],strict=True):
        for field in ('cnpj','documento','status','diferencas'):
            assert row[field]==target[field], (row,target)
        assert len(row['sped'])==target['sped']
        assert len(row['dominio'])==target['dominio']
    (examples/'resultado_simulado.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    xlsx=post('/exportar/xlsx'); sheet=load_workbook(BytesIO(xlsx)).active
    assert sheet.max_row==106
    assert Counter(sheet.cell(i,1).value for i in range(2,107))==Counter(expected['summary'])
    (examples/'resultado_simulado.xlsx').write_bytes(xlsx)
    report=post('/exportar/pdf')
    with pdfplumber.open(BytesIO(report)) as pdf:
        pages=len(pdf.pages)
        assert pages>1
        text='\n'.join(p.extract_text() or '' for p in pdf.pages)
        for row in expected['rows']:
            assert row['documento'] in text
    (examples/'resultado_simulado.pdf').write_bytes(report)
    for source in ('sped','dominio'):
        with opener.open(f'http://127.0.0.1:5000/exemplos/{source}_simulado.pdf',timeout=10) as response:
            assert response.read()==(examples/f'{source}_simulado.pdf').read_bytes()
    summary=dict(resultado='APROVADO',registros=result['counts'],chaves=105,situacoes=expected['summary'],avisos=0,linhas_excel=106,paginas_relatorio_pdf=pages,validacao='HTTP real e comparação das 105 chaves com gabarito independente; exportações reabertas e verificadas. Não é homologação com arquivos reais.')
    (examples/'validacao.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
