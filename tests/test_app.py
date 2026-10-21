from io import BytesIO
from decimal import Decimal
from reportlab.pdfgen import canvas
from openpyxl import load_workbook
from app import app
from comparator import compare, parse_text

def pdf(lines):
    stream=BytesIO(); c=canvas.Canvas(stream)
    for i,line in enumerate(lines): c.drawString(30,800-i*20,line)
    c.save(); stream.seek(0); return stream

def payload():
    return {'sped':(pdf(['12.345.678/0001-90 | 001 | Empresa A | 1.250,00','12.345.678/0001-90 | 2 | Empresa A | 50,00']), 'sped.pdf'),'dominio':(pdf(['12345678000190 | 1 | Empresa A | 1.251,00','12345678000190 | 3 | Empresa A | 50,00']), 'dominio.pdf')}

def test_full_pdf_comparison_and_exports():
    client=app.test_client()
    assert client.get('/').status_code==200
    response=client.post('/comparar',data=payload())
    assert response.status_code==200
    data=response.json
    assert data['counts']==[2,2]
    assert [r['status'] for r in data['rows']]==['Divergente','Somente SPED','Somente Domínio']
    assert data['rows'][0]['diferencas']=='Valor'
    xlsx=client.post('/exportar/xlsx',data=payload())
    wb=load_workbook(BytesIO(xlsx.data))
    assert wb.active.max_row==4
    report=client.post('/exportar/pdf',data=payload())
    assert report.status_code==200 and report.data.startswith(b'%PDF-')

def test_duplicates_and_supplier():
    a,_=parse_text('12345678000190 | 01 | Ação LTDA | 0,10',1)
    b,_=parse_text('12345678000190 | 1 | ACAO LTDA | 0,10',1)
    assert compare(a,b)[0]['status']=='Conferente'
    assert compare(a+a,b)[0]['status']=='Duplicidade / revisar'
    b[0]['fornecedor']='Outro'
    assert compare(a,b)[0]['diferencas']=='Fornecedor'

def test_reject_unreadable_and_invalid():
    client=app.test_client()
    assert client.post('/comparar',data={}).status_code==422
    data=payload();data['sped']=(BytesIO(b'bad'),'a.pdf')
    assert client.post('/comparar',data=data).status_code==422
    data=payload();data['sped']=(pdf(['Relatorio sem notas reconhecidas']), 'a.pdf')
    assert client.post('/comparar',data=data).status_code==422

def test_no_silent_invalid_row():
    notes,issues=parse_text('12345678000190 | 1 | Empresa | 2.30',1)
    assert not notes and issues

def test_formula_is_exported_as_text():
    data=payload();data['sped']=(pdf(['12345678000190 | 1 | =1+1 | 1,00']), 'a.pdf')
    response=app.test_client().post('/exportar/xlsx',data=data)
    wb=load_workbook(BytesIO(response.data))
    assert wb.active['E2'].data_type=='s'
