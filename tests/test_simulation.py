from collections import Counter
from io import BytesIO
import pdfplumber
from openpyxl import load_workbook
from app import app
from scripts.gerar_exemplos import generate


def test_paginated_simulation_against_independent_manifest(tmp_path):
    expected = generate(tmp_path)
    def payload():
        return {name:(BytesIO((tmp_path/f'{name}_simulado.pdf').read_bytes()), f'{name}.pdf') for name in ('sped','dominio')}
    for source in ('sped','dominio'):
        with pdfplumber.open(tmp_path/f'{source}_simulado.pdf') as pdf:
            assert len(pdf.pages) >= 4
    client=app.test_client()
    response=client.post('/comparar',data=payload())
    assert response.status_code==200
    actual=response.json
    assert actual['counts']==expected['counts']==[100,95]
    assert actual['warnings']==[]
    assert dict(Counter(r['status'] for r in actual['rows']))==expected['summary']
    assert len(actual['rows'])==105
    for row, target in zip(actual['rows'],expected['rows'], strict=True):
        assert {k:row[k] for k in ('cnpj','documento','status','diferencas')}=={k:target[k] for k in ('cnpj','documento','status','diferencas')}
        assert len(row['sped'])==target['sped']
        assert len(row['dominio'])==target['dominio']
    response=client.post('/exportar/xlsx',data=payload())
    assert response.status_code==200
    ws=load_workbook(BytesIO(response.data)).active
    assert ws.max_row==106
    assert Counter(ws.cell(i,1).value for i in range(2,107))==Counter(expected['summary'])
    response=client.post('/exportar/pdf',data=payload())
    assert response.status_code==200
    with pdfplumber.open(BytesIO(response.data)) as pdf:
        assert len(pdf.pages)>1
        texts='\n'.join(p.extract_text() or '' for p in pdf.pages)
        for r in expected['rows']:
            assert r['documento'] in texts
        assert 'Duplicidade' in texts and 'Somente SPED' in texts and 'Somente Domínio' in texts


def test_example_downloads():
    client = app.test_client()
    for name in ('sped_simulado.pdf', 'dominio_simulado.pdf'):
        response = client.get('/exemplos/' + name)
        assert response.status_code == 200
        assert response.data.startswith(b'%PDF-')
        assert 'attachment' in response.headers['Content-Disposition']
    assert client.get('/exemplos/requirements.txt').status_code == 404
