from copy import deepcopy
from io import BytesIO

import pdfplumber
import pytest
from openpyxl import load_workbook

import app as application
from comparator import compare, supplier_matches


@pytest.mark.parametrize('left,right', [
    ('Empresa Exemplo LTDA.', 'EMPRESA EXEMPLO LTDA'),
    ('Ação, Comércio LTDA', 'acao comercio ltda.'),
    ('Exemplo S.A.', 'Exemplo SA'),
    ('Empresa Exemplo LTDA', 'Empresa Exemplo'),
    ('Empresa Exemplo LIMITADA', 'Empresa Exemplo LTDA.'),
    ('Empresa Exemplo Limitada', 'Empresa Exemplo'),
    ('Empresa Exemplo LTDA (LOJA)', 'Empresa Exemplo'),
    ('Empresa Exemplo LTDA (LOJA EXEMPLO)', 'Empresa Exemplo LTDA'),
    ('Empresa Exemplo LTDA', 'Empresa Exemplo LTDA (LOJA EXEMPLO)'),
    ('Empresa Exemplo LTDA (LOJA) (CENTRO)', 'Empresa Exemplo LTDA.'),
    ('Empresa Exemplo', 'Empresa Exemplo Comercio LTDA'),
    ('Empresa Exemplo SA', 'Empresa Exemplo'),
    ('Empresa Exemplo 1', 'Empresa Exemplo 2'),
    ('Exempl LTDA (LOJA)', 'Exemplo LTDA'),
    ('Empresa Exemplo (CENTRO)', 'Empresa Exemplo (BAIRRO)'),
    ('Agro Horizonte Comercio de Sementes e Produtos Agricolas LTDA', 'Agro Horizonte Comercio'),
    ('Agro Horizonte Produtos Agricolas', 'Agro Horizonte Comercio e Representacoes'),
    ('Agro Horizonte Comercio de Sementes', 'Agro Horizonte Comercio de Seme'),
    ('Agroexemplo Comercio de Produtos', 'Agroexemplo Distribuidora'),
    ('Joa Batista Silva LTDA (LOJA)', 'Joao Batista Silva LTDA'),
    ('12.345.678 MARCOS ANTONIO EXEMPLO', 'MARCOS ANTONIO EXEMPLO'),
    ('MARCA COMERCIO AURORA DIGITAL LTDA', 'COMERCIO AURORA DIGITAL LTDA'),
    ('D A ALVES COSTA LTDA', 'ALVES COSTA E COSTA LTDA- ME'),
])
def test_supplier_formatting_and_extra_trade_name(left, right):
    assert supplier_matches(left, right)
    assert supplier_matches(right, left)


@pytest.mark.parametrize('left,right', [
    ('Empresa Um LTDA', 'Empresa Dois LTDA'),
    ('LTDA Exemplo', 'Exemplo'),
    ('LTDA', 'LIMITADA'),
    ('Empresa (LOJA)', 'Empresa Outra'),
    ('Empresa (CENTRO) LTDA', 'Empresa LTDA'),
    ('Comercio de Sementes Aurora', 'Comercio de Pecas Horizonte'),
    ('Comercio e Industria Alfa', 'Comercio e Industria Beta'),
    ('Comercio de Produtos Alfa', 'Comercio de Produtos Beta'),
    ('Transportes Alfa', 'Transportes Beta'),
    ('Alfa', 'Alfabeto'),
    ('(EMPRESA)', 'EMPRESA DIFERENTE'),
    ('...', ','),
    ('', ''),
])
def test_meaningful_name_changes_still_require_review(left, right):
    assert not supplier_matches(left, right)
    assert not supplier_matches(right, left)


def notes():
    base = dict(cnpj='12345678000190', documento='123', pagina=1, valor='340.00')
    return ([dict(base, fornecedor='Empresa Exemplo LTDA (LOJA EXEMPLO)')],
            [dict(base, fornecedor='Empresa Exemplo LTDA', codigo_fornecedor='0042')])


def test_internal_code_is_metadata_and_values_still_checked():
    left, right = notes()
    originals = deepcopy((left, right))
    assert compare(left, right)[0]['status'] == 'Conferente'
    assert (left, right) == originals  # Preserve the original names for review.
    right[0]['codigo_fornecedor'] = '9999'
    assert compare(left, right)[0]['status'] == 'Conferente'
    right[0]['valor'] = '339.99'
    row = compare(left, right)[0]
    assert row['status'] == 'Divergente' and row['diferencas'] == 'Valor'
    right[0]['fornecedor'] = ''
    right[0]['valor'] = '340.00'
    assert compare(left, right)[0]['status'] == 'Conferência parcial'


def test_exports_separate_code_from_preserved_name(monkeypatch):
    monkeypatch.setattr(application, 'analyze', lambda: dict(rows=compare(*notes()), warnings=[]))
    client = application.app.test_client()
    response = client.post('/exportar/xlsx')
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.data)).active
    assert sheet['A2'].value == 'Conferente'
    assert sheet['E2'].value == 'Empresa Exemplo LTDA (LOJA EXEMPLO)'
    assert sheet['G2'].value == 'Empresa Exemplo LTDA\nCódigo interno: 0042'
    assert sheet['G2'].alignment.wrap_text
    response = client.post('/exportar/pdf')
    assert response.status_code == 200
    with pdfplumber.open(BytesIO(response.data)) as pdf:
        text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
    assert 'Conferente' in text and 'Código interno: 0042' in text


def test_omitted_legal_suffix_still_checks_identity_document_and_amount():
    left, right = notes()
    left[0]['fornecedor'] = 'Empresa Exemplo LTDA'
    right[0]['fornecedor'] = 'Empresa Exemplo'
    assert compare(left, right)[0]['status'] == 'Conferente'
    right[0]['valor'] = '340.01'
    assert compare(left, right)[0]['diferencas'] == 'Valor'
    right[0]['cnpj'] = '98765432000110'
    assert {r['status'] for r in compare(left, right)} == {'Somente SPED', 'Somente Domínio'}
    right[0]['cnpj'] = left[0]['cnpj']
    right[0]['documento'] = '124'
    assert {r['status'] for r in compare(left, right)} == {'Somente SPED', 'Somente Domínio'}


def test_export_identifies_supplier_name_difference(monkeypatch):
    left, right = notes()
    right[0]['fornecedor'] = 'Outra Empresa'
    monkeypatch.setattr(application, 'analyze', lambda: dict(rows=compare(left, right), warnings=[]))
    response = application.app.test_client().post('/exportar/xlsx')
    sheet = load_workbook(BytesIO(response.data)).active
    assert sheet['D2'].value == 'Nome do fornecedor'
    assert 'Código interno: 0042' in sheet['G2'].value


def test_long_name_never_overrides_money_or_identity_and_name_is_last():
    left, right = notes()
    left[0]['fornecedor'] = 'Agro Horizonte Comercio de Sementes e Produtos LTDA'
    right[0]['fornecedor'] = 'Agro Horizonte Comercio'
    assert compare(left, right)[0]['status'] == 'Conferente'
    right[0]['valor'] = '339.99'
    assert compare(left, right)[0]['diferencas'] == 'Valor'
    right[0]['fornecedor'] = 'Transportadora Planalto'
    assert compare(left, right)[0]['diferencas'] == 'Valor, Fornecedor'
    right[0]['cnpj'] = '98765432000110'
    assert {r['status'] for r in compare(left, right)} == {'Somente SPED', 'Somente Domínio'}
    right[0]['cnpj'] = left[0]['cnpj']
    right[0]['documento'] = '124'
    assert {r['status'] for r in compare(left, right)} == {'Somente SPED', 'Somente Domínio'}
