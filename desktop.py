"""Launcher for the offline Windows installation; no pip or external server."""
from io import BytesIO
import argparse
import json
import os
from pathlib import Path
import threading
import urllib.request
import uuid
import webbrowser

from app import app, APP_VERSION
from waitress import create_server

ROOT = Path(__file__).resolve().parent


def local_opener():
    # Local PDFs stay on the computer, even when a corporate proxy is set.
    return urllib.request.build_opener(urllib.request.ProxyHandler({}))


def wait_for_instance(address, instance_id):
    import time
    for _ in range(60):
        try:
            with local_opener().open(address + '/versao', timeout=1) as response:
                if json.load(response).get('instancia') == instance_id:
                    return
        except (OSError, ValueError):
            pass
        time.sleep(.25)
    raise RuntimeError('O comparador não ficou pronto para abrir.')


def multipart_examples():
    boundary = 'comparador-' + uuid.uuid4().hex
    chunks = []
    for source in ('sped', 'dominio'):
        filename = source + '_simulado.pdf'
        chunks.append((f'--{boundary}\r\nContent-Disposition: form-data; name="{source}"; '
                       f'filename="{filename}"\r\nContent-Type: application/pdf\r\n\r\n').encode())
        chunks.extend([(ROOT / 'exemplos' / filename).read_bytes(), b'\r\n'])
    chunks.append(f'--{boundary}--\r\n'.encode())
    return b''.join(chunks), 'multipart/form-data; boundary=' + boundary


def self_test(address, instance_id):
    """Exercise installed native PDF libraries, HTTP and both report formats."""
    from collections import Counter
    import pdfplumber
    from openpyxl import load_workbook
    wait_for_instance(address, instance_id)
    opener = local_opener()
    with opener.open(address, timeout=10) as response:
        assert f'Versão {APP_VERSION}' in response.read().decode()
    for filename in ('sped_simulado.pdf', 'dominio_simulado.pdf'):
        with opener.open(address + '/exemplos/' + filename, timeout=10) as response:
            assert response.read().startswith(b'%PDF-')
    payload, content_type = multipart_examples()
    manifest = json.loads((ROOT / 'exemplos' / 'gabarito.json').read_text(encoding='utf-8'))

    def post(path):
        req = urllib.request.Request(address + path, data=payload, headers={'Content-Type': content_type})
        with opener.open(req, timeout=90) as response:
            return response.read()

    result = json.loads(post('/comparar'))
    assert result['counts'] == manifest['counts']
    assert result['warnings'] == []
    assert Counter(row['status'] for row in result['rows']) == manifest['summary']
    assert len(result['rows']) == len(manifest['rows'])
    for actual, expected in zip(result['rows'], manifest['rows'], strict=True):
        assert all(actual[key] == expected[key] for key in ('cnpj', 'documento', 'status', 'diferencas'))
    sheet = load_workbook(BytesIO(post('/exportar/xlsx'))).active
    assert sheet.max_row == len(result['rows']) + 1
    with pdfplumber.open(BytesIO(post('/exportar/pdf'))) as pdf:
        assert len(pdf.pages) > 1
        assert 'Comparativo' in (pdf.pages[0].extract_text() or '')
    print(json.dumps({'versao': APP_VERSION, 'self_test': 'ok', 'documentos': len(result['rows']),
                      'counts': result['counts'], 'relatorios': ['xlsx', 'pdf']}, ensure_ascii=False), flush=True)


def main(argv=None):
    parser = argparse.ArgumentParser(description='Comparador de notas Lucrattiva')
    parser.add_argument('--self-test', action='store_true', help='Validar instalação com PDFs fictícios, sem abrir o navegador.')
    args = parser.parse_args(argv)
    instance_id = uuid.uuid4().hex
    os.environ['COMPARADOR_INSTANCE_ID'] = instance_id
    # Bind port 0 atomically: occupied ports and old copies cannot be selected.
    server = create_server(app, host='127.0.0.1', port=0, threads=4)
    address = f'http://127.0.0.1:{server.effective_port}'
    print(f'Comparador de Notas Lucrattiva — versão {APP_VERSION}', flush=True)
    print('Endereço desta cópia: ' + address, flush=True)
    try:
        if args.self_test:
            thread = threading.Thread(target=server.run, daemon=True)
            thread.start()
            self_test(address, instance_id)
        else:
            def open_browser():
                try:
                    wait_for_instance(address, instance_id)
                    if not webbrowser.open(address):
                        print('Abra no navegador o endereço desta cópia mostrado acima.', flush=True)
                except Exception as exc:
                    print(f'Não foi possível abrir o navegador: {exc}', flush=True)
            print('Mantenha esta janela aberta enquanto usar. Feche-a ou pressione Ctrl+C para encerrar.', flush=True)
            threading.Thread(target=open_browser, daemon=True).start()
            server.run()
    except KeyboardInterrupt:
        pass
    finally:
        server.close()
        server.task_dispatcher.shutdown()
    return 0


if __name__ == '__main__':
    try:
        raise SystemExit(main())
    except Exception:
        import traceback
        traceback.print_exc()
        if os.name == 'nt':
            import ctypes
            ctypes.windll.user32.MessageBoxW(None, 'Não foi possível abrir o comparador. Confira o erro na janela do programa ou reinstale a versão mais recente.', 'Comparador de Notas', 0x10)
        raise SystemExit(1)
